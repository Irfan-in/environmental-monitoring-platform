import os
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

SRC_PATH = r"C:\Users\Irfan\Downloads\S3 MCA SCRUM BOOK (1).docx"
OUT_PATH_DOWNLOADS = r"C:\Users\Irfan\Downloads\S3_MCA_SCRUM_BOOK_Filled.docx"
OUT_PATH_PROJECT = r"c:\project\Irfan\S3_MCA_SCRUM_BOOK_Filled.docx"

doc = docx.Document(SRC_PATH)

def set_font(run, name="Times New Roman", size_pt=11, bold=False, italic=False, color=None):
    run.font.name = name
    if size_pt:
        run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color

# ==========================================
# 1. COVER PAGE UPDATES
# ==========================================
# P3: 20MCA245 – MINI PROJECT(3RD SEM MCA 2024-27) -> 20MCA245 – MINI PROJECT(3RD SEM MCA 2025-27)
p3 = doc.paragraphs[3]
for r in p3.runs:
    if "2024-27" in r.text:
        r.text = r.text.replace("2024-27", "2025-27")

# P6: NAME
p6 = doc.paragraphs[6]
p6.text = ""
r_space = p6.add_run("                                         ")
set_font(r_space, "Times New Roman", 15, bold=True)
r_lbl = p6.add_run("NAME:    ")
set_font(r_lbl, "Times New Roman", 15, bold=True)
r_val = p6.add_run("Mohammed Irfan KS")
set_font(r_val, "Times New Roman", 15, bold=True)

# P7: Class
p7 = doc.paragraphs[7]
p7.text = ""
r_space = p7.add_run("                                         ")
set_font(r_space, "Times New Roman", 15, bold=True)
r_lbl = p7.add_run("Class:        ")
set_font(r_lbl, "Times New Roman", 15, bold=True)
r_val = p7.add_run("MCA 3rd Sem")
set_font(r_val, "Times New Roman", 15, bold=True)

# P8: Reg.No
p8 = doc.paragraphs[8]
p8.text = ""
r_space = p8.add_run("                                         ")
set_font(r_space, "Times New Roman", 15, bold=True)
r_lbl = p8.add_run("Reg.No:     ")
set_font(r_lbl, "Times New Roman", 15, bold=True)
r_val = p8.add_run("KCMIT25MCA-2011")
set_font(r_val, "Times New Roman", 15, bold=True)

# ==========================================
# 2. META TABLE (Table 0)
# ==========================================
t0 = doc.tables[0]

# Row 0: Name of the Student
c0 = t0.rows[0].cells[1]
c0.text = ""
p = c0.paragraphs[0]
r = p.add_run("Mohammed Irfan KS")
set_font(r, "Times New Roman", 12, bold=True)

# Row 1: Batch
c1 = t0.rows[1].cells[1]
c1.text = ""
p = c1.paragraphs[0]
r = p.add_run("2025 - 2027")
set_font(r, "Times New Roman", 12, bold=True)

# Row 2: Register Number
t0.rows[2].cells[0].text = "Register Number:"
p_lbl = t0.rows[2].cells[0].paragraphs[0]
set_font(p_lbl.runs[0], "Times New Roman", 13, bold=True)

c2 = t0.rows[2].cells[1]
c2.text = ""
p = c2.paragraphs[0]
r = p.add_run("KCMIT25MCA-2011")
set_font(r, "Times New Roman", 12, bold=True)

# Row 3: Academic Year (leave blank as requested)
c3 = t0.rows[3].cells[1]
c3.text = ""

# Row 4: Subject Name & Code (already 20MCA245 - MINI PROJECT)
# Keep as is

# Row 5: Title of the Project
c5 = t0.rows[5].cells[1]
c5.text = ""
p = c5.paragraphs[0]
r = p.add_run("Integrated Environmental Monitoring Platform: IoT Microclimate Sensing, Cloud Ingestion, and Real-Time Cross-Platform Telemetry")
set_font(r, "Times New Roman", 11, bold=True)

# Row 6: Name and Designation of the Guide (leave blank as requested)
c6 = t0.rows[6].cells[1]
c6.text = ""

# ==========================================
# 3. TOPIC & MODULE DESCRIPTION
# ==========================================
# P70 is "MINI PROJECT TOPIC DESCRIPTION AND MODULE DESCRIPTION"
# P71 through P95 are empty paragraphs before P96 ("WEEKLY REPORT OF THE PROJECT")

topic_content = [
    ("PROJECT TITLE: INTEGRATED ENVIRONMENTAL MONITORING PLATFORM", "heading", 12, True),
    ("", "empty", 0, False),
    ("1. PROJECT OVERVIEW & PROBLEM STATEMENT", "subheading", 11, True),
    ("The Integrated Environmental Monitoring Platform is an end-to-end, multi-tiered IoT and mobile telemetry platform designed to monitor, analyze, and visualize real-time indoor microclimate conditions alongside localized outdoor meteorological parameters. Maintaining optimal indoor air temperature and relative humidity is critical in sensitive settings such as academic server rooms, laboratories, and modern workplaces to prevent hardware degradation, control humidity risks, and ensure occupant thermal comfort. The platform bridges physical embedded hardware sensing with scalable cloud aggregation and cross-platform mobile decision support.", "body", 10.5, False),
    ("", "empty", 0, False),
    ("2. MODULAR ARCHITECTURE BREAKDOWN", "subheading", 11, True),
    ("The project is engineered across four core modular subsystems:", "body", 10.5, False),
    ("• Module 1: Edge Sensing Node & Hardware Firmware (C++ / Arduino IDE / FreeRTOS): Interfaces an ESP32 dual-core 240MHz microcontroller with a DHT22 high-accuracy digital sensor (sampling temperature at ±0.5°C and relative humidity at ±2% RH) and an SSD1306 0.96-inch monochrome I2C OLED display. Employs non-blocking timers, diagnostic visual metrics, and resilient Wi-Fi state management with automated reconnect routines to dispatch JSON telemetry via HTTP POST requests.", "bullet", 10, False),
    ("• Module 2: Cloud Ingestion Gateway & Meteorological Sync (Python / FastAPI / SQLite): High-performance RESTful API service engineered with FastAPI and Uvicorn, hosted on Render cloud infrastructure. Manages a relational SQLite database with SQLAlchemy ORM modeling for persistent time-series indexing. Orchestrates background synchronization with the Open-Meteo external weather API to fetch local ambient outdoor meteorological parameters (temperature, humidity, wind speed, precipitation probability).", "bullet", 10, False),
    ("• Module 3: Comparative Analytics & Thermal Comfort Engine (ASHRAE 55 / Steadman Algorithm): Real-time computational engine calculating the Steadman Heat Index ('apparent temperature') to evaluate physiological heat stress. Derives indoor vs. outdoor microclimate differentials to assess building insulation efficiency. Assesses thermal comfort compliance against ASHRAE Standard 55 envelopes, outputting 'Optimal', 'Acceptable', or 'Discomfort' ratings with actionable ventilation guidance and deduplicated threshold alerting.", "bullet", 10, False),
    ("• Module 4: Mobile Client & Local-First Storage (Dart / Flutter / Provider): Cross-platform mobile client for Android built with Flutter 3, featuring Provider state management. Delivers real-time telemetry cards and dual-trace interactive comparative line charts (fl_chart) with UTC-to-local IST (+5:30) synchronization. Integrates LocalStorageService for persistent on-device JSON caching, enabling offline operation, alongside Role-Based Access Control (RBAC) supporting guest preview and authenticated user sessions.", "bullet", 10, False),
    ("", "empty", 0, False),
    ("3. HARDWARE & SOFTWARE SPECIFICATIONS", "subheading", 11, True),
    ("• Microcontroller: ESP32 NodeMCU-32S (Xtensa Dual-Core 32-bit LX6 @ 240 MHz, 520 KB SRAM, 2.4 GHz Wi-Fi)", "spec", 10, False),
    ("• Sensors & Display: DHT22 (AM2302) Digital Sensor, SSD1306 0.96\" 128x64 Monochrome OLED Display (I2C 0x3C)", "spec", 10, False),
    ("• Cloud Backend & Database: Python 3.12, FastAPI, Uvicorn, SQLAlchemy ORM, SQLite, Render Cloud Hosting", "spec", 10, False),
    ("• Mobile Framework: Google Flutter 3.x, Dart SDK, Provider State Management, fl_chart, SharedPreferences", "spec", 10, False),
    ("• External APIs: Open-Meteo REST Weather API for localized reference meteorology", "spec", 10, False),
]

# We populate existing paragraphs P71 onward
current_p_idx = 71
for text, kind, size, bold in topic_content:
    if current_p_idx < 96:
        p = doc.paragraphs[current_p_idx]
        p.text = ""
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        if kind == "heading":
            r = p.add_run(text)
            set_font(r, "Times New Roman", size, bold=True)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(2)
        elif kind == "subheading":
            r = p.add_run(text)
            set_font(r, "Times New Roman", size, bold=True)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(2)
        elif kind == "bullet" or kind == "spec":
            r = p.add_run(text)
            set_font(r, "Times New Roman", size, bold=False)
            p.paragraph_format.left_indent = Inches(0.2)
        elif kind == "body":
            r = p.add_run(text)
            set_font(r, "Times New Roman", size, bold=False)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        elif kind == "empty":
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
        current_p_idx += 1

# If there are remaining empty paragraphs between current_p_idx and 96, remove their text/spacing
while current_p_idx < 96:
    p = doc.paragraphs[current_p_idx]
    p.text = ""
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    current_p_idx += 1

# ==========================================
# 4. WEEKLY PROGRESS REPORTS (Tables 7, 8, 9, 10)
# ==========================================
weekly_data = [
    {
        "table_idx": 7,
        "week": "Week: 1",
        "dates": "From: 15-09-2026        To: 21-09-2026",
        "work_done": [
            "• Conducted domain literature survey on smart environmental monitoring systems, ASHRAE thermal standards, and IoT microclimate tracking architectures.",
            "• Formulated system requirements and procured hardware components: ESP32 NodeMCU-32S, DHT22 (AM2302) digital temperature/humidity sensor, and SSD1306 0.96\" I2C OLED display.",
            "• Prototyped hardware schematic on breadboard, configured 3.3V power rails, and verified I2C bus device discovery at address 0x3C.",
            "• Implemented baseline ESP32 firmware in C++/Arduino IDE utilizing Adafruit DHT and SSD1306 libraries for continuous ambient sampling.",
            "• Conducted initial calibration and validation of temperature and relative humidity readouts displayed on the local OLED screen."
        ],
        "remark": "Hardware circuit design and local sensor display verified. Finalize cloud database schema and REST API specifications for sprint 2."
    },
    {
        "table_idx": 8,
        "week": "Week: 2",
        "dates": "From: 22-09-2026        To: 28-09-2026",
        "work_done": [
            "• Designed and implemented cloud RESTful backend using Python FastAPI and Uvicorn ASGI framework.",
            "• Modeled SQLite relational database schema utilizing SQLAlchemy ORM for structured time-series telemetry storage and audit indexing.",
            "• Integrated Open-Meteo external meteorological API to synchronously fetch localized outdoor ambient parameters based on geographic coordinates.",
            "• Enhanced ESP32 firmware with non-blocking WiFiClient and HTTPClient routines to dispatch JSON telemetry payloads via HTTP POST requests.",
            "• Validated end-to-end edge-to-cloud data transmission, payload integrity, and database insertion via Swagger/OpenAPI documentation."
        ],
        "remark": "Backend data ingestion endpoints and database models approved. Proceed to cross-platform mobile dashboard development."
    },
    {
        "table_idx": 9,
        "week": "Week: 3",
        "dates": "From: 29-09-2026        To: 06-10-2026",
        "work_done": [
            "• Scaffolded cross-platform mobile client in Flutter 3, employing the Provider pattern for reactive state management.",
            "• Engineered responsive dashboard UI with status badges, temperature/humidity metric cards, and microclimate indicators.",
            "• Integrated interactive dual-trace line charts using fl_chart to visualize synchronized indoor telemetry versus outdoor weather curves.",
            "• Implemented mathematical formulas for Steadman's Heat Index and ASHRAE Standard 55 indoor thermal comfort compliance estimation.",
            "• Designed intelligent, deduplicated threshold alert system to detect extreme humidity conditions and heat stress risks."
        ],
        "remark": "Mobile dashboard UI and comparative telemetry charts validated. Implement persistent storage, handle edge cases, and deploy to cloud."
    },
    {
        "table_idx": 10,
        "week": "Week: 4",
        "dates": "From: 07-10-2026        To: 18-10-2026",
        "work_done": [
            "• Deployed production FastAPI backend to Render cloud platform (https://environmental-monitoring-backend.onrender.com).",
            "• Implemented LocalStorageService in Flutter for on-device persistent JSON telemetry caching, enabling seamless offline operation.",
            "• Diagnosed and resolved UTC-to-local IST (+5:30) timestamp desynchronization across interactive charts and telemetry cards.",
            "• Implemented Role-Based Access Control (RBAC) with secure session handling, supporting guest preview and authenticated user profiles (User: Irfan).",
            "• Compiled release APK, performed physical device testing on Samsung Galaxy A55 (SM-A556E), and validated end-to-end telemetry pipeline."
        ],
        "remark": "Comprehensive end-to-end integration verified across embedded, cloud, and mobile tiers. Project documentation and Scrum Book approved for final viva."
    }
]

for report in weekly_data:
    t = doc.tables[report["table_idx"]]
    
    # Row 0: Dates
    c0 = t.rows[0].cells[0]
    p_date = c0.paragraphs[2]
    p_date.text = report["dates"]
    set_font(p_date.runs[0], "Times New Roman", 11, bold=True)
    
    # Row 1: Work Done
    c1 = t.rows[1].cells[0]
    # Keep the first paragraph (header)
    p_hdr = c1.paragraphs[0]
    p_hdr.text = "1. Work Done (Should be written by students before meeting your guides)"
    set_font(p_hdr.runs[0], "Times New Roman", 11, bold=True)
    p_hdr.paragraph_format.space_before = Pt(3)
    p_hdr.paragraph_format.space_after = Pt(3)
    
    # Add each work item as a clean paragraph
    for item in report["work_done"]:
        p_item = c1.add_paragraph()
        p_item.paragraph_format.space_before = Pt(1)
        p_item.paragraph_format.space_after = Pt(1)
        p_item.paragraph_format.line_spacing = 1.15
        p_item.paragraph_format.left_indent = Inches(0.15)
        r = p_item.add_run(item)
        set_font(r, "Times New Roman", 10, bold=False)
        
    # Row 2: Suggestions/Remark by Guide
    c2 = t.rows[2].cells[0]
    c2.text = ""
    p_rem_hdr = c2.paragraphs[0]
    r_hdr = p_rem_hdr.add_run("2. Suggestions/Remark by Guide:")
    set_font(r_hdr, "Times New Roman", 11, bold=True)
    p_rem_hdr.paragraph_format.space_before = Pt(3)
    p_rem_hdr.paragraph_format.space_after = Pt(2)
    
    p_rem_val = c2.add_paragraph()
    p_rem_val.paragraph_format.space_before = Pt(1)
    p_rem_val.paragraph_format.space_after = Pt(3)
    p_rem_val.paragraph_format.line_spacing = 1.15
    p_rem_val.paragraph_format.left_indent = Inches(0.15)
    r_val = p_rem_val.add_run(report["remark"])
    set_font(r_val, "Times New Roman", 10.5, italic=True)

    # Row 3: Initials by Guide
    c3 = t.rows[3].cells[0]
    # Keep original 4. Initials by : GUIDE

# Save to output locations
doc.save(OUT_PATH_DOWNLOADS)
doc.save(OUT_PATH_PROJECT)
print("SUCCESS: Scrum Book successfully generated at:")
print(" -", OUT_PATH_DOWNLOADS)
print(" -", OUT_PATH_PROJECT)
