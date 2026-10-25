import os
import docx
from docx.shared import Pt

SRC_PATH = r"C:\Users\Irfan\Downloads\Smart Vehicle Scrumbook.docx"
OUT_PATHS = [
    r"C:\Users\Irfan\Downloads\S3_MCA_SCRUM_BOOK_15Weeks.docx",
    r"c:\project\Irfan\S3_MCA_SCRUM_BOOK_15Weeks.docx",
    r"C:\Users\Irfan\Downloads\S3_MCA_SCRUM_BOOK_Filled.docx",
    r"c:\project\Irfan\S3_MCA_SCRUM_BOOK_Filled.docx"
]

doc = docx.Document(SRC_PATH)

def set_font(run, name="Times New Roman", size_pt=14, bold=True, italic=False):
    run.font.name = name
    if size_pt:
        run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic

# ==========================================
# 1. COVER PAGE UPDATES (Page 1)
# ==========================================
# P8: NAME : MEGHANATH T -> NAME : Mohammed Irfan KS
p8 = doc.paragraphs[8]
p8.text = ""
r_space = p8.add_run("             ")
set_font(r_space, "Times New Roman", 16, bold=True)
r_lbl = p8.add_run("NAME :    ")
set_font(r_lbl, "Times New Roman", 16, bold=True)
r_val = p8.add_run("Mohammed Irfan KS")
set_font(r_val, "Times New Roman", 16, bold=True)

# P9: Class : S3 MCA
p9 = doc.paragraphs[9]
p9.text = ""
r_space = p9.add_run("    ")
set_font(r_space, "Times New Roman", 16, bold=True)
r_lbl = p9.add_run("Class    :    ")
set_font(r_lbl, "Times New Roman", 16, bold=True)
r_val = p9.add_run("S3 MCA")
set_font(r_val, "Times New Roman", 16, bold=True)

# P10: Reg.No: KCMIT25MCA-2010 -> Reg.No: KCMIT25MCA-2011
p10 = doc.paragraphs[10]
p10.text = ""
r_lbl = p10.add_run("Reg.No:     ")
set_font(r_lbl, "Times New Roman", 16, bold=True)
r_val = p10.add_run("KCMIT25MCA-2011")
set_font(r_val, "Times New Roman", 16, bold=True)

# ==========================================
# 2. META TABLE (Table 0, Page 2)
# ==========================================
t0 = doc.tables[0]

# Row 0: Name of the Student
c0 = t0.rows[0].cells[1]
c0.text = ""
p = c0.paragraphs[0]
r = p.add_run("  Mohammed Irfan KS")
set_font(r, "Times New Roman", 14, bold=True)

# Row 1: Batch (2025-2027)
c1 = t0.rows[1].cells[1]
c1.text = ""
p = c1.paragraphs[0]
r = p.add_run("2025-2027")
set_font(r, "Times New Roman", 14, bold=True)

# Row 2: Register Number
c2 = t0.rows[2].cells[1]
c2.text = ""
p = c2.paragraphs[0]
r = p.add_run("  KCMIT25MCA-2011")
set_font(r, "Times New Roman", 14, bold=True)

# Row 3: Academic Year (leave blank as requested by user)
c3 = t0.rows[3].cells[1]
c3.text = ""

# Row 4: Subject Name & Code
c4 = t0.rows[4].cells[1]
c4.text = ""
p = c4.paragraphs[0]
r = p.add_run(" 20MCA245 - MINI PROJECT")
set_font(r, "Times New Roman", 14, bold=True)

# Row 5: Title of the Project
c5 = t0.rows[5].cells[1]
c5.text = ""
p = c5.paragraphs[0]
r = p.add_run("  INTEGRATED ENVIRONMENTAL MONITORING PLATFORM")
set_font(r, "Times New Roman", 12, bold=True)

# Row 6: Name and Designation of the Guide (leave blank as requested by user)
c6 = t0.rows[6].cells[1]
c6.text = ""

# ==========================================
# 3. TOPIC & MODULE DESCRIPTION (Page 6)
# ==========================================
p55 = doc.paragraphs[55]
p55.text = "MINI PROJECT TOPIC DESCRIPTION"
set_font(p55.runs[0], "Times New Roman", 14, bold=True)

p56 = doc.paragraphs[56]
p56.text = "Project Title: Integrated Environmental Monitoring Platform"
set_font(p56.runs[0], "Times New Roman", 14, bold=True)

p57 = doc.paragraphs[57]
p57.text = "Project Description:"
set_font(p57.runs[0], "Times New Roman", 14, bold=False)

p58 = doc.paragraphs[58]
p58.text = "Integrated Environmental Monitoring Platform is an IoT-based system designed to monitor and analyze indoor microclimate conditions alongside localized outdoor weather data. It uses an ESP32 microcontroller with a DHT22 sensor to capture environmental metrics and transmits them to a cloud-hosted FastAPI backend. The system provides real-time dashboards, dual-trace comparative analytics, ASHRAE 55 thermal comfort scoring, offline persistence, and role-based user access to support climate monitoring."
set_font(p58.runs[0], "Times New Roman", 14, bold=False)

p59 = doc.paragraphs[59]
p59.text = "MODULE DESCRIPTION"
set_font(p59.runs[0], "Times New Roman", 14, bold=False)

modules_data = [
    (60, 61, "1. Microclimate Sensing and Hardware Firmware",
     "This module interfaces the ESP32 microcontroller with a DHT22 sensor and OLED display to collect temperature and humidity readings and transmit telemetry over Wi-Fi."),
    (62, 63, "2. Cloud Ingestion and Weather Sync",
     "This module uses FastAPI to receive sensor data, store records in an SQLite database, and retrieve localized outdoor weather from the Open-Meteo API."),
    (64, 65, "3. Comparative Analytics and Comfort Engine",
     "This module calculates Steadman's Heat Index and evaluates indoor thermal comfort using ASHRAE 55 standards to provide actionable environmental insights."),
    (66, 67, "4. Mobile Dashboard and Visualization",
     "This module provides an interactive Flutter mobile interface with real-time metric cards, synchronized comparative charts, and environmental alert notifications."),
    (68, 69, "5. Local Storage and User Management",
     "This module enables offline data caching using LocalStorageService and provides role-based authentication supporting user profiles and guest access.")
]

for title_idx, desc_idx, title_text, desc_text in modules_data:
    p_t = doc.paragraphs[title_idx]
    p_t.text = title_text
    set_font(p_t.runs[0], "Times New Roman", 14, bold=False)
    
    p_d = doc.paragraphs[desc_idx]
    p_d.text = desc_text
    set_font(p_d.runs[0], "Times New Roman", 14, bold=False)

# ==========================================
# 4. 15 WEEKLY PROGRESS REPORTS (Tables 7 to 21)
# ==========================================
weekly_reports = [
    {
        "table_idx": 7,
        "week_no": 1,
        "dates": "From:06/07/2026     To: 10/07/2026.",
        "subtitle": "Project Initiation and Requirement Analysis",
        "work_done": "The project was initiated by identifying the problem statement and operational requirements of an Integrated Environmental Monitoring Platform. The core system capabilities, including real-time indoor microclimate sensing, cloud-based data ingestion, localized outdoor meteorological synchronization, and cross-platform mobile telemetry visualization, were outlined. The multi-tiered architecture was planned, selecting the ESP32 microcontroller and DHT22 digital sensor for hardware sensing, Python FastAPI for the backend REST API, SQLite for structured storage, and Flutter for the mobile client.",
        "remark": "Proceed with hardware component selection, circuit schematics, and development environment setup."
    },
    {
        "table_idx": 8,
        "week_no": 2,
        "dates": "From:13/07/2026     To: 17/07/2026\t.",
        "subtitle": "Hardware Selection and Prototyping Setup",
        "work_done": "The hardware component specifications were finalized and components were procured, including an ESP32 NodeMCU-32S microcontroller, DHT22 (AM2302) digital temperature and relative humidity sensor, and a 0.96-inch monochrome SSD1306 I2C OLED display. The circuit schematic was designed and prototyped on a breadboard, verifying regulated 3.3V power rails, pull-up resistors for the DHT22 single-bus data line, and I2C SDA/SCL pin assignments (GPIO 21 and 22). Baseline firmware tools were configured in the Arduino IDE with necessary board definition packages.",
        "remark": "Verify sensor accuracy, I2C address resolution, and establish stable pinout connections."
    },
    {
        "table_idx": 9,
        "week_no": 3,
        "dates": "From:20/07/2026     To: 24/07/2026\t.",
        "subtitle": "Sensor Interfacing and Local Display Implementation",
        "work_done": "Firmware was developed in C++ to interface the ESP32 with the DHT22 sensor and SSD1306 OLED display using Adafruit DHT and SSD1306 libraries. An I2C scanner routine was executed to confirm display communication at address 0x3C. The firmware logic was implemented to sample temperature and humidity values at non-blocking 2-second intervals, format the readings, and render real-time indoor environmental metrics on the OLED screen. Preliminary bench testing verified reading consistency and display responsiveness.",
        "remark": "Implement non-blocking sampling intervals and prepare Wi-Fi communication routines."
    },
    {
        "table_idx": 10,
        "week_no": 4,
        "dates": "From:27/07/2026    To: 31/07/2026",
        "subtitle": "ESP32 Wi-Fi Integration and Telemetry Serialization",
        "work_done": "The edge firmware was expanded to include Wi-Fi networking using the ESP32 WiFi library. An automated connection handler was implemented to authenticate against local access points and manage automated reconnection routines upon network interruption. The ArduinoJson library was integrated to structure temperature, relative humidity, computed heat index, and device status into serialized JSON telemetry payloads. Local network HTTP client routines were established to prepare for cloud transmission.",
        "remark": "Design the backend REST API specifications and database models for telemetry ingestion."
    },
    {
        "table_idx": 11,
        "week_no": 5,
        "dates": "From.:03/08/2026   To: 07/08/2026",
        "subtitle": "Backend Architecture and Database Design",
        "work_done": "The cloud backend structure was developed using Python FastAPI and the Uvicorn ASGI framework. The relational database schema was modeled using SQLAlchemy ORM to manage telemetry tables, capturing sensor timestamps, temperature, relative humidity, computed heat index, and hardware status flags. An SQLite database was initialized with structured time-series indexing. Pydantic data schemas were defined to enforce strict request validation, payload typing, and robust error handling for incoming IoT telemetry endpoints.",
        "remark": "Proceed with API endpoint development and test telemetry ingestion with simulated requests."
    },
    {
        "table_idx": 12,
        "week_no": 6,
        "dates": "From.:10/08/2026    To: 14/08/2026",
        "subtitle": "Ingestion API Development and Testing",
        "work_done": "RESTful API endpoints for telemetry ingestion (/api/telemetry) and retrieval (/api/telemetry/latest, /api/telemetry/history) were implemented in FastAPI. The endpoints were tested using Swagger UI and Postman to ensure reliable JSON parsing, automated UTC timestamp generation, and efficient database insertion. Data sanitization routines were integrated to validate sensor value ranges and filter abnormal readings. Query response times for time-series records were tested and optimized.",
        "remark": "Integrate the external meteorological weather service to provide outdoor reference data."
    },
    {
        "table_idx": 13,
        "week_no": 7,
        "dates": "From.:17/08/2026   To: 21/08/2026",
        "subtitle": "External Meteorological API Integration",
        "work_done": "An external meteorological synchronization pipeline was developed in the FastAPI backend utilizing the Open-Meteo REST API. A background task was created to query localized outdoor meteorological data based on geographic coordinates, capturing ambient temperature, relative humidity, wind speed, and precipitation probability. The backend logic was arranged to align outdoor reference data alongside indoor sensor telemetry, enabling synchronized dual-source environmental analysis.",
        "remark": "Develop the comparative analytics algorithms and thermal comfort evaluation models."
    },
    {
        "table_idx": 14,
        "week_no": 8,
        "dates": "From.:24/08/2026      To: 28/08/2026",
        "subtitle": "Thermal Comfort and Comparative Analytics Engine",
        "work_done": "Analytical computation modules were developed to derive meaningful environmental insights from indoor and outdoor datasets. The Steadman Heat Index formula was implemented to compute apparent (\"feels like\") temperature accounting for humidity-driven evaporative cooling reduction. Thermal comfort boundaries were modeled based on ASHRAE Standard 55 envelopes, classifying indoor conditions into \"Optimal\", \"Acceptable\", or \"Discomfort\" states. Indoor-outdoor differentials were calculated to evaluate building thermal retention and HVAC efficiency.",
        "remark": "Initialize the mobile application project and design the primary user interface layout."
    },
    {
        "table_idx": 15,
        "week_no": 9,
        "dates": "From.:31/08/2026      To: 04/09/2026",
        "subtitle": "Mobile Application Setup and Architecture",
        "work_done": "The cross-platform mobile client application was created using Google Flutter 3 and Dart SDK. The application architecture was established using the Provider pattern to enable scalable state management and reactive UI updates. Core project folders for screens, models, services, providers, and widgets were structured. Theme data featuring clean modern typography and responsive card layouts was configured, and basic screen routing was set up and tested on an Android emulator.",
        "remark": "Implement HTTP network service and bind live telemetry data to dashboard cards."
    },
    {
        "table_idx": 16,
        "week_no": 10,
        "dates": "From.:07/09/2026     To: 11/09/2026",
        "subtitle": "Mobile Dashboard and Telemetry Cards",
        "work_done": "The primary dashboard interface was developed in Flutter to present live environmental telemetry. Responsive metric cards were implemented to display real-time indoor temperature, relative humidity, heat index, and thermal comfort category status. An HTTP client service was integrated to fetch telemetry updates from the backend REST endpoints. Visual indicators, color-coded badges, and icons were arranged to visually represent environmental quality and status levels.",
        "remark": "Develop interactive graphical charts to visualize comparative telemetry trends over time."
    },
    {
        "table_idx": 17,
        "week_no": 11,
        "dates": "From.:14/09/2026      To: 18/09/2026",
        "subtitle": "Comparative Chart Visualization Development",
        "work_done": "Interactive time-series graphical charts were integrated into the Flutter application utilizing the fl_chart library. Dual-trace line graphs were implemented to render indoor sensor telemetry directly alongside localized outdoor Open-Meteo meteorological curves. Interactive tooltips, touch markers, and customized grid lines were configured. Timestamp parsing routines were developed to convert UTC cloud records into local IST (UTC+5:30) time to ensure accurate chart timeline alignment.",
        "remark": "Implement on-device local storage persistence to enable offline access and historical auditing."
    },
    {
        "table_idx": 18,
        "week_no": 12,
        "dates": "From.:21/09/2026     To: 25/09/2026",
        "subtitle": "Local Storage and Offline Data Persistence",
        "work_done": "A local persistence layer was developed using LocalStorageService and SharedPreferences in Flutter. The service was engineered to cache recent telemetry records in persistent JSON format on device storage upon receipt. Logic was implemented to check network connectivity, allowing the mobile application to seamlessly load cached historical records, view past trends, and maintain dashboard functionality when disconnected from Wi-Fi or backend services.",
        "remark": "Implement user authentication, role-based access control, and climate alert notifications."
    },
    {
        "table_idx": 19,
        "week_no": 13,
        "dates": "From.:28/09/2026       To: 1/10/2026",
        "subtitle": "User Management and Alert Engine",
        "work_done": "Role-Based Access Control (RBAC) was implemented in the Flutter application, supporting authenticated user profiles (\"User: Irfan\") and an unrestricted guest mode for open telemetry monitoring. A session management service was integrated to preserve login state across app restarts. Additionally, an intelligent environmental alert system was created to monitor humidity thresholds and thermal comfort boundaries, displaying single relevant alerts for current climate conditions while preventing alert duplication.",
        "remark": "Deploy the backend to cloud infrastructure and conduct end-to-end integration testing."
    },
    {
        "table_idx": 20,
        "week_no": 14,
        "dates": "From.:05/10/2026      To:09/10/2026",
        "subtitle": "Cloud Deployment and End-to-End Edge-Cloud Testing",
        "work_done": "The FastAPI backend service was deployed to Render cloud hosting (https://environmental-monitoring-backend.onrender.com) with persistent database storage. The physical ESP32 edge node was programmed with production cloud URLs and tested under live conditions, reliably transmitting DHT22 sensor readings over Wi-Fi. End-to-end data transmission, cloud ingestion, Open-Meteo weather synchronization, and mobile dashboard updates were validated across sustained continuous operating cycles.",
        "remark": "Perform physical device verification, final optimizations, and prepare presentation artifacts."
    },
    {
        "table_idx": 21,
        "week_no": 15,
        "dates": "From.:12/10/2026      To:16/10/2026",
        "subtitle": "Physical Device Testing, Refinement and Final Verification",
        "work_done": "The mobile application was compiled into a release APK and installed on a physical smartphone (Samsung Galaxy A55 5G) for real-world field testing. UI responsiveness, app icon rendering, real-time graph synchronization, and offline cache retrieval were verified under actual operating conditions. Minor display and font styling adjustments were refined. The entire system architecture was finalized, and the mini-project documentation and Scrum Book were compiled for evaluation and viva voce.",
        "remark": "The project is complete, fully validated, and ready for demonstration. Documentation approved for final submission."
    }
]

for rep in weekly_reports:
    t = doc.tables[rep["table_idx"]]
    
    # Row 1: Subtitle & Work Done
    c1 = t.rows[1].cells[0]
    p_sub = c1.paragraphs[1]
    p_sub.text = rep["subtitle"]
    set_font(p_sub.runs[0], "Times New Roman", 12, bold=True)
    
    p_work = c1.paragraphs[2]
    p_work.text = rep["work_done"]
    set_font(p_work.runs[0], "Times New Roman", 11.5, bold=False)
    
    # Row 2: Guide Remark
    c2 = t.rows[2].cells[0]
    p_rem = c2.paragraphs[1]
    p_rem.text = rep["remark"]
    set_font(p_rem.runs[0], "Times New Roman", 11.5, bold=False)

for pth in OUT_PATHS:
    doc.save(pth)
    print("Saved successfully to:", pth)
